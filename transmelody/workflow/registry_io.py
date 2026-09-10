"""Portable XLSX registry storage; no app-specific JavaScript runtime required."""
from pathlib import Path
import os
import uuid

FIELDS = ('id', 'title', 'original_filename', 'numbered_filename',
          'source_folder', 'added_date', 'status', 'notes')
HEADERS = ('编号', '曲目名', '原始文件名', '编号后文件名', '来源文件夹', '登记日期', '状态', '备注')


def validate(rows):
    ids = [int(row['id']) for row in rows]
    if any(i < 1 for i in ids) or len(ids) != len(set(ids)):
        raise ValueError('记录表编号必须是唯一的正整数。')
    return ids


def read(path):
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('记录表必须是 .xlsx 文件。')
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        values = wb.worksheets[0].iter_rows(max_col=8, values_only=True)
        if tuple(next(values, ())) != HEADERS:
            raise ValueError('记录表表头不符合预期。')
        rows = [{k: '' if v is None else str(v).strip() for k, v in zip(FIELDS, row)}
                for row in values if row[0] is not None and str(row[0]).strip()]
        validate(rows)
        return rows
    finally:
        wb.close()


def write(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    path = Path(path)
    ids = validate(rows)
    wb = Workbook()
    sheet = wb.active
    sheet.title = '曲目记录'
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([int(row['id']), *[str(row.get(k, '')) for k in FIELDS[1:]]])
    # Titles/notes are literal text, never Excel formulas supplied by filenames.
    for cells in sheet.iter_rows(min_row=2, min_col=2, max_col=8):
        for cell in cells:
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    for cell in sheet[1][:8]:
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.font = Font(color='FFFFFF', bold=True)
    for i, width in enumerate((8, 28, 28, 20, 18, 14, 16, 28), 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:H{max(1, len(rows) + 1)}'
    for address, value in {'J1': '摘要', 'K1': '数值', 'J2': '已登记曲目',
                           'K2': len(ids), 'J3': '下一个编号', 'K3': max(ids, default=0) + 1}.items():
        sheet[address] = value
    sheet.column_dimensions['J'].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f'.registry-{uuid.uuid4().hex}.xlsx')
    try:
        wb.save(tmp)
        assert read(tmp) == [{k: str(int(r['id'])) if k == 'id' else str(r.get(k, '')).strip()
                              for k in FIELDS} for r in rows]
        os.replace(tmp, path)
    finally:
        wb.close()
        tmp.unlink(missing_ok=True)


def append(path, new_rows):
    write(path, read(path) + new_rows)


def update(path, updates):
    rows = read(path)
    by_id = {r['id']: r for r in rows}
    for change in updates:
        row = by_id[str(change['id'])]
        for field in ('status', 'notes'):
            if field in change:
                row[field] = str(change[field])
    write(path, rows)
